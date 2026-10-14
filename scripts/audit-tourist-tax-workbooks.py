from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook
import json, re

FILES = [
    Path(r'C:\Users\kosti\Downloads\1_База_міста_та_ставки_377.xlsx'),
    Path(r'C:\Users\kosti\Downloads\2_Готовий_SEO_контент_377.xlsx'),
]

def read(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(x).strip() if x is not None else '' for x in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(v is not None for v in row)]

def normalized(s):
    return re.sub(r'\s+', ' ', str(s or '').strip().lower())

def pct(text):
    if text is None: return []
    return [float(x.replace(',', '.')) for x in re.findall(r'(?<![\d.])(\d+(?:[,.]\d+)?)\s*%', str(text))]

base, seo = map(read, FILES)
by_slug = lambda rows: {normalized(r.get('SEO Slug (URL)') or r.get('slug')): r for r in rows}
bmap, smap = by_slug(base), by_slug(seo)
duplicates = {k:v for k,v in Counter(normalized(r.get('SEO Slug (URL)') or r.get('slug')) for r in base+seo).items() if v>2}
missing = {'in_seo_only': sorted(set(smap)-set(bmap)), 'in_base_only': sorted(set(bmap)-set(smap))}
rate_conflicts = defaultdict(set)
community_names = {}
for r in base:
    # Community names are not globally unique. Keep the source workbook's
    # region dimension in the audit key so equal labels in different oblasts
    # cannot merge or hide rate conflicts.
    region = normalized(r.get('Область'))
    community = normalized(r.get('Територіальна громада (ТГ)'))
    c = f'{region}|{community}'
    community_names[c] = {'region': r.get('Область'), 'community': r.get('Територіальна громада (ТГ)')}
    rate_conflicts[c].add((tuple(pct(r.get('Ставка (Внутрішній турист)'))), tuple(pct(r.get("Ставка (В'їзний/Іноземний)")))))
per_page = []
for slug in sorted(set(bmap)&set(smap)):
    b,s=bmap[slug],smap[slug]
    br=(tuple(pct(b.get('Ставка (Внутрішній турист)'))),tuple(pct(b.get("Ставка (В'їзний/Іноземний)"))))
    sr=[]
    for key in ('Block1 direct intro','Block2 calculator params/formula','Block3 exemption','Block4 hotels/agents','Block5 FAQ','Meta Description','Meta Title','H1'):
        sr += pct(s.get(key))
    used_8000 = any('8 000' in str(s.get(k) or '') or '8000' in str(s.get(k) or '') for k in s)
    if used_8000 or (br[0] and any(x not in sr for x in br[0])) or (br[1] and any(x not in sr for x in br[1])):
        per_page.append({'slug':slug,'locality':b.get('Населений пункт (SEO)'),'rates_base':br,'seo_percents':sorted(set(sr)),'has_8000':used_8000})
legal_patterns = {
  'broad_child_exemption':r'дітей[^.]{0,80}(?:до 18|звільн|не є платник)',
  'broad_military_exemption':r'(?:ветеран|учасник бойових|військов)[^.]{0,120}(?:звільн|не є платник|пільг)',
  'broad_vpo_exemption':r'(?:ВПО|внутрішньо переміщ)[^.]{0,120}(?:звільн|не є платник|пільг)',
  'automatic_deadline_10days':r'10\s*(?:календарних|робочих)?\s*дн',
  'automatic_funds_usage':r'(?:кошти|збір)[^.]{0,100}(?:спрям|використов).{0,80}(?:туризм|інфраструктур)',
  'required_document_claim':r'(?:свідоцтв[оа]|паспорт)[^.]{0,100}(?:обов|потріб|необхід)',
}
legal_hits=[]
for r in seo:
    body=' '.join(str(v) for k,v in r.items() if any(t in k.lower() for t in ['block','faq']))
    hits=[name for name,pattern in legal_patterns.items() if re.search(pattern,body,re.I)]
    if hits: legal_hits.append({'slug':r.get('slug'),'city':r.get('locality'),'flags':hits})
jsonld=[]
for r in seo:
    raw=r.get('Block6 FAQPage JSON-LD')
    if not raw: continue
    try:
        parsed=json.loads(str(raw))
        if not isinstance(parsed,dict) or parsed.get('@type')!='FAQPage': raise ValueError('not FAQPage object')
    except Exception as e: jsonld.append({'slug':r.get('slug'),'error':str(e)})
summary={
 'file_rows':[len(base),len(seo)],
 'base_headers':list(base[0]) if base else [],
 'seo_headers':list(seo[0]) if seo else [],
 'slug_counts':[len(bmap),len(smap)],'slug_mismatches':missing,
 'community_count':len(rate_conflicts),
 'community_rate_conflict_count':sum(len(v)>1 for v in rate_conflicts.values()),
 'community_rate_conflicts':[{'region':community_names[k]['region'],'community':community_names[k]['community'],'rate_pairs':list(v)} for k,v in rate_conflicts.items() if len(v)>1][:30],
 'page_rate_or_base_flags_count':len(per_page),'page_rate_or_base_flags_sample':per_page[:30],
 'legal_review_flags_count':len(legal_hits),'legal_review_flags_sample':legal_hits[:30],
 'faq_jsonld_errors_count':len(jsonld),'faq_jsonld_errors_sample':jsonld[:15],
 'source_decision_url_columns':[k for k in base[0] if re.search(r'(джерело|рішення|url|період|статус|katottg)',k,re.I)],
}
print(json.dumps(summary,ensure_ascii=False,indent=2,default=list))
