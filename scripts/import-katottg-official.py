"""Join VDALO tourist-tax workbook places to the pinned official KATOTTG release."""
from collections import Counter, defaultdict
from pathlib import Path
import csv, hashlib, json, re, unicodedata
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'seo-workspace/tourist-tax/sources/kodifikator-2026-07-07.xlsx'
WORKBOOK=ROOT/'versions/v1/tourist-tax-workbook-data.json'
OUTPUT=ROOT/'versions/v1/tourist-tax-katottg-2026-07-07.json'
REVIEW=ROOT/'seo-workspace/tourist-tax/exports/katottg-match-review.csv'
SOURCE_URL='https://mininfra.gov.ua/storage/app/sites/1/uploaded-files/kodifikator-07-07.xlsx'
RELEASE_DATE='2026-07-07'

# Exact locality-code reconciliations for same-name and renamed locations.
# Codes are taken from the pinned official KATOTTG workbook and the intended
# community is cross-checked before any override is accepted.
MANUAL_PLACE_MATCHES={
    'bar': {'placeCode':'UA05060010010059821','communityCode':'UA05060010000053179','note':'Офіційний КАТОТТГ містить у Барській громаді місто Бар і однойменне селище; обрано місто категорії M за точним кодом громади з бази.'},
    'pivdenne': {'placeCode':'UA51100410010044384','communityCode':'UA51100410000059549','note':'Назву громади в базі звірено з чинною Південнівською громадою за офіційним КАТОТТГ; місто Південне відрізнено від однойменного села за кодом громади.'},
    'yuzhnoukrayinsk': {'placeCode':'UA48040250010045068','communityCode':'UA48040250000044166','note':'У базі використано попередню назву Южноукраїнськ; офіційний КАТОТТГ фіксує чинну назву Південноукраїнськ у цій громаді.'},
}

def norm(value):
    value=unicodedata.normalize('NFKC',str(value or '')).casefold().replace('’',"'").replace('ʼ',"'")
    value=re.sub(r'\b(область|області|обл\.?|район|району|р-н)\b',' ',value)
    value=re.sub(r'[^\w\s-]',' ',value,flags=re.UNICODE)
    return re.sub(r'\s+',' ',value).strip(' .')

def norm_community(value):
    value=norm(value)
    value=re.sub(r'\b(територіальна громада|міська|селищна|сільська|міської|селищної|сільської|тг)\b',' ',value)
    return re.sub(r'\s+',' ',value).strip()

def load_official():
    wb=load_workbook(SOURCE,read_only=True,data_only=True)
    ws=wb['Кодифікатор']
    items=[]
    for row in ws.iter_rows(min_row=5,values_only=True):
        region,district,community,place,additional,category,name=row
        if not category or not name: continue
        items.append({'regionCode':region,'districtCode':district,'communityCode':community,
          'placeCode':place,'additionalCode':additional,'category':str(category).strip(),
          'name':str(name).strip()})
    districts={item['districtCode']:item['name'] for item in items if item['category']=='P' and item['districtCode']}
    return items,districts

def main():
    data=json.loads(WORKBOOK.read_text(encoding='utf-8'))
    official,district_names=load_official()
    regions={norm(x['name']):x['regionCode'] for x in official if x['category']=='O'}
    kyiv=next((x for x in official if x['category']=='K' and norm(x['name'])=='київ'),None)
    sevastopol=next((x for x in official if x['category']=='K' and norm(x['name'])=='севастополь'),None)
    communities=defaultdict(list)
    places=defaultdict(list)
    for item in official:
        if item['category']=='H': communities[(item['regionCode'],norm(item['name']))].append(item)
        if item['category'] in ('M','C','X'): places[(item['regionCode'],norm(item['name']))].append(item)
    out=[]
    counters=Counter()
    with REVIEW.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f)
        writer.writerow(['slug','place_name','workbook_region','workbook_community','official_place_name','place_category','settlement_katottg','official_community_name','community_katottg','district_katottg','region_katottg','match_status','match_basis','review_note'])
        for source in data['places']:
            region_name=source.get('region','')
            region_code=regions.get(norm(region_name))
            is_kyiv=norm(region_name) in ('м київ','київ місто','місто київ')
            if is_kyiv and kyiv: region_code=kyiv['regionCode']
            if norm(region_name) in ('м севастополь','севастополь місто') and sevastopol: region_code=sevastopol['regionCode']
            comm_key=norm_community(source.get('community'))
            comm_candidates=communities.get((region_code,comm_key),[]) if region_code else []
            name_key=norm(source.get('name'))
            place_candidates=places.get((region_code,name_key),[]) if region_code else []
            basis='exact_locality_and_region'
            manual_note=''
            duplicate_locality_slug=None
            # Some source slugs append a region, attraction, or administrative
            # qualifier in parentheses. Drop it only when the base locality and
            # a uniquely matched official parent community identify one code.
            parenthetical=re.search(r'\s*\([^)]*\)\s*$',str(source.get('name') or ''))
            if not place_candidates and parenthetical and len(comm_candidates)==1:
                base_name=str(source.get('name') or '')[:parenthetical.start()].strip()
                base_candidates=places.get((region_code,norm(base_name)),[]) if region_code else []
                community_code=comm_candidates[0]['communityCode']
                by_comm=[candidate for candidate in base_candidates if candidate['communityCode']==community_code]
                if len(by_comm)==1:
                    prior=next((item for item in out if item.get('settlementCode')==by_comm[0]['placeCode']),None)
                    if prior:
                        duplicate_locality_slug=prior['slug']
                    else:
                        place_candidates=by_comm
                        basis='exact_base_locality_region_and_community_after_qualifier_removal'
                        manual_note=f"Офіційний код підтверджено за назвою «{base_name}», областю та єдиною громадою з КАТОТТГ; дужковий уточнювач «{str(source.get('name'))[parenthetical.start():].strip()}» використано лише як редакційний опис."
                        counters['qualified_locality_match']+=1
            manual_match=MANUAL_PLACE_MATCHES.get(source['slug'])
            if manual_match:
                exact=next((item for item in official if item['placeCode']==manual_match['placeCode'] and item['category'] in ('M','C','X')),None)
                if not exact or exact['regionCode']!=region_code or exact['communityCode']!=manual_match['communityCode']:
                    raise ValueError(f"Manual KATOTTG override no longer matches pinned source: {source['slug']}")
                parent=next((item for item in official if item['category']=='H' and item['communityCode']==manual_match['communityCode']),None)
                if not parent:
                    raise ValueError(f"Manual KATOTTG override has no exact community: {source['slug']}")
                place_candidates=[exact]
                comm_candidates=[parent]
                basis='reviewed_exact_official_code_override'
                manual_note=manual_match['note']
            # When duplicate locality names exist, an exact territorial-community
            # match can safely disambiguate them if it leaves one KATOTTG object.
            if len(place_candidates)>1 and comm_candidates:
                comm_codes={candidate['communityCode'] for candidate in comm_candidates}
                by_comm=[candidate for candidate in place_candidates if candidate['communityCode'] in comm_codes]
                if len(by_comm)==1:
                    place_candidates=by_comm
                    basis='exact_locality_region_and_community'
            place=place_candidates[0] if len(place_candidates)==1 else None
            # Kyiv itself is a first-level KATOTTG object, not an O/H/M hierarchy.
            if not place and is_kyiv and kyiv and name_key=='київ':
                place=kyiv
                comm_candidates=[kyiv]
                basis='exact_special_city_name_and_code'
            if place and place['category']=='M' and len(comm_candidates)!=1:
                # The parent community recorded by KATOTTG is stronger than a
                # workbook spelling variant of an old community name.
                comm_candidates=[x for x in official if x['category']=='H' and x['communityCode']==place['communityCode']]
            community=comm_candidates[0] if len(comm_candidates)==1 else None
            community_code=place['communityCode'] if place else (community['communityCode'] if community else None)
            # A row in category K represents its own community and locality.
            if place and place['category']=='K':
                community_code=place['regionCode']
                community=place
            if duplicate_locality_slug:
                status='LOCALITY_ALIAS_DUPLICATE'
                note=f"Після зняття редакційного уточнювача населений пункт збігається з уже зіставленим записом {duplicate_locality_slug}; той самий код не дублюється."
                counters['place_alias_duplicates']+=1
            elif place:
                status='MATCHED'
                note=manual_note
                counters['place_matched']+=1
                if manual_match: counters['manual_code_override']+=1
            elif not region_code:
                status='REGION_UNMATCHED'
                note='Workbook region is not an exact first-level KATOTTG match.'
                counters['region_unmatched']+=1
            elif len(place_candidates)>1:
                status='AMBIGUOUS_LOCALITY'
                note='More than one official locality has this name in the region; confirm the exact object.'
                counters['place_ambiguous']+=1
            else:
                status='LOCALITY_NOT_IN_KATOTTG'
                note='Workbook name is not an exact KATOTTG locality name; it may be a destination alias or require manual research.'
                counters['place_unmatched']+=1
            if community_code: counters['community_code_matched']+=1
            else: counters['community_code_missing']+=1
            settlement_code=(place.get('regionCode') if place and place['category']=='K' else place.get('placeCode')) if place else None
            record={
              'slug':source['slug'],'settlementCode':settlement_code,
              'officialSettlementName':place.get('name') if place else None,
              'settlementCategory':place.get('category') if place else None,
              'regionCode':region_code,'officialRegionName':next((x['name'] for x in official if x['category']=='O' and x['regionCode']==region_code),kyiv['name'] if is_kyiv and kyiv else None),
              'districtCode':place.get('districtCode') if place else None,
              'officialDistrictName':district_names.get(place.get('districtCode')) if place else None,
              'communityCode':community_code,'officialCommunityName':community.get('name') if community else None,
              'status':status,'matchBasis':basis if place else None,'reviewNote':note,
            }
            out.append(record)
            writer.writerow([source['slug'],source.get('name'),region_name,source.get('community'),record['officialSettlementName'],record['settlementCategory'],record['settlementCode'],record['officialCommunityName'],record['communityCode'],record['districtCode'],record['regionCode'],status,record['matchBasis'],note])
    codes=[item['settlementCode'] for item in out if item['settlementCode']]
    if len(codes)!=len(set(codes)):
        duplicates=[code for code,n in Counter(codes).items() if n>1]
        raise SystemExit(f'Duplicate settlement codes assigned: {duplicates[:10]}')
    payload={'version':'katottg-2026-07-07-v1','releaseDate':RELEASE_DATE,'officialSourceUrl':SOURCE_URL,
      'sourceSha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'matchPolicy':'Exact normalized locality and oblast names; duplicate locality names require a unique matching community; unmatched and ambiguous values remain null.',
      'categoryCodes':{'O':'oblast','P':'district','H':'territorial community','M':'city','C':'village','X':'settlement','K':'special city'},
      'manualCodeOverrides':MANUAL_PLACE_MATCHES,
      'summary':dict(counters),'places':out}
    OUTPUT.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'output':str(OUTPUT),'reviewCsv':str(REVIEW),'sourceSha256':payload['sourceSha256'],'officialRows':len(official),'placeRows':len(out),'summary':dict(counters)},ensure_ascii=True,indent=2))

if __name__=='__main__':main()
