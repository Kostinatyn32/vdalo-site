"""Reproducible, non-publishing import of the owner-supplied tourist-tax books."""
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook
import json, re, hashlib

ROOT=Path(__file__).resolve().parents[1]
INPUTS=[
 Path(r'C:\Users\kosti\Downloads\1_База_міста_та_ставки_377.xlsx'),
 Path(r'C:\Users\kosti\Downloads\2_Готовий_SEO_контент_377.xlsx'),
]
OUTPUT=ROOT/'versions/v1/tourist-tax-workbook-data.json'

def read(path):
    book=load_workbook(path,read_only=True,data_only=True)
    sheet=book[book.sheetnames[0]]
    rows=list(sheet.iter_rows(values_only=True))
    names=[str(v).strip() if v is not None else '' for v in rows[0]]
    return [dict(zip(names,row)) for row in rows[1:] if any(v is not None for v in row)]

def slugify(value):
    return re.sub(r'[^a-z0-9-]+','-',str(value or '').strip().lower()).strip('-')

def number(value):
    if value is None:return None
    try:return float(str(value).replace(' ','').replace(',','.'))
    except ValueError:return None

def rate(value):
    found=re.findall(r'(?<![\d.])(\d+(?:[,.]\d+)?)\s*%',str(value or ''))
    return float(found[0].replace(',','.')) if found else None

def legacy_amount(rate_value):
    return round(8000*rate_value/100,2) if rate_value is not None else None

base,seo=map(read,INPUTS)
bmap={slugify(row.get('SEO Slug (URL)')):row for row in base}
smap={slugify(row.get('SEO Slug (URL)') or row.get('slug')):row for row in seo}
base_cols=list(base[0])
seo_cols=list(seo[0])
if len(base)!=377 or len(seo)!=377:raise SystemExit(f'Expected 377 rows in each workbook, got {len(base)} and {len(seo)}')
if len(bmap)!=len(base) or len(smap)!=len(seo):raise SystemExit('Duplicate or empty slug detected')
if set(bmap)!=set(smap):raise SystemExit('Workbook slugs do not reconcile')

places=[]
communities=defaultdict(list)
for slug,b in sorted(bmap.items()):
    s=smap[slug]
    # The supplied workbooks have stable column positions but contain headers
    # with inconsistent spelling, so resolve content by documented position.
    rates={'domestic':rate(b[base_cols[7]]),'inbound':rate(b[base_cols[9]])}
    community=str(b[base_cols[5]] or '').strip()
    item={
      'slug':slug,'route':f'/tourist-tax/{slug}/',
      'name':str(b[base_cols[0]] or '').strip(),
      'cluster':str(b[base_cols[2]] or '').strip(),
      'category':str(b[base_cols[3]] or '').strip(),
      'region':str(b[base_cols[4]] or '').strip(),
      'community':community,
      'council':str(b[base_cols[6]] or '').strip(),
      'proposedRates':rates,
      'legacyAmountsBasedOn8000':{
        'domestic':legacy_amount(rates['domestic']),
        'inbound':legacy_amount(rates['inbound']),
      },
      'budgetCode':str(b[base_cols[11]] or '').strip(),
      'localNotes':str(b[base_cols[12]] or '').strip(),
      'preparedSeo':{
        'title':str(s[seo_cols[5]] or '').strip(),
        'description':str(s[seo_cols[6]] or '').strip(),
        'h1':str(s[seo_cols[7]] or '').strip(),
        'intro':str(s[seo_cols[8]] or '').strip(),
        'calculator':str(s[seo_cols[9]] or '').strip(),
        'exemptions':str(s[seo_cols[10]] or '').strip(),
        'agents':str(s[seo_cols[11]] or '').strip(),
        'faq':str(s[seo_cols[12]] or '').strip(),
        'faqJsonLd':str(s[seo_cols[13]] or '').strip(),
      },
      'review':{'geography':'REVIEW_REQUIRED','localDecision':'SECONDARY_SOURCE_ONLY','seoLegalClaims':'REVIEW_REQUIRED','indexable':False},
      'sourceFiles':[p.name for p in INPUTS],
    }
    places.append(item);communities[community].append(item)

community_rows=[]
for name,members in sorted(communities.items()):
    proposed=sorted({(m['proposedRates']['domestic'],m['proposedRates']['inbound']) for m in members})
    community_rows.append({
      'name':name,'council':members[0]['council'],'placeCount':len(members),
      'katottgCode':None,'decisionNumber':None,'decisionDate':None,'effectiveFrom':None,'effectiveTo':None,'sourceUrl':None,
      'proposedRates':{'domestic':proposed[0][0] if len(proposed)==1 else None,'inbound':proposed[0][1] if len(proposed)==1 else None},
      'rateConflict':len(proposed)>1,'rateVariants':[{'domestic':a,'inbound':b} for a,b in proposed],
      'status':'SECONDARY_SOURCE_ONLY','checkedAt':None,
    })

payload={
 'schemaVersion':1,'importedAt':'2026-10-14',
 'sourceFiles':[{'name':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'rows':len(rows)} for p,rows in zip(INPUTS,(base,seo))],
 'counts':{'places':len(places),'communities':len(community_rows),'rateConflictCommunities':sum(c['rateConflict'] for c in community_rows)},
 'policy':{'sourceRatesAreNotVerified':True,'legacyAmountBase':8000,'2026AmountBase':8647,'doNotCalculateFromProposedRates':True,'doNotIndexUntilAllReadinessGatesPass':True},
 'communities':community_rows,'places':places,
}
OUTPUT.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'output':str(OUTPUT),'counts':payload['counts'],'sourceFiles':payload['sourceFiles']},ensure_ascii=True,indent=2))
